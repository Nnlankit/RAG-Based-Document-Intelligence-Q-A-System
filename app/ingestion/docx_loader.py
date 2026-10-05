"""DOCX document extractor using python-docx."""

import logging
from pathlib import Path
from typing import List, Optional
import docx

from app.core.exceptions import DocumentCorruptedError
from app.ingestion.cleaner import TextCleaner
from app.ingestion.loader import ExtractedPage

logger = logging.getLogger("document_intelligence.ingestion.docx")


class DocxLoader:
    """Extracts structured text, headings, and tables from DOCX documents."""

    @classmethod
    def extract(
        cls,
        file_path: Path,
        document_id: str,
        filename: str,
    ) -> List[ExtractedPage]:
        """Extracts paragraphs grouped by heading/section.
        
        Args:
            file_path: Path to the .docx file.
            document_id: Unique document ID.
            filename: Original file name.
            
        Returns:
            List of ExtractedPage objects with page_number=None and heading section attribution.
        """
        try:
            doc = docx.Document(str(file_path))
        except Exception as e:
            logger.error(f"Failed to open DOCX '{filename}': {e}", exc_info=True)
            raise DocumentCorruptedError(f"Failed to open or parse corrupted DOCX file: {filename}")

        sections: List[ExtractedPage] = []
        current_section = "Introduction"
        current_paragraphs: List[str] = []

        def flush_current_section():
            if current_paragraphs:
                full_text = "\n\n".join(current_paragraphs)
                cleaned = TextCleaner.clean(full_text)
                if cleaned:
                    sections.append(
                        ExtractedPage(
                            text=cleaned,
                            document_id=document_id,
                            filename=filename,
                            page_number=None,  # Section 12 rule: do not invent page numbers for DOCX
                            source_type="docx",
                            section=current_section,
                        )
                    )
                current_paragraphs.clear()

        # 1. Process paragraphs and headings
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name if p.style else ""
            if "Heading" in style_name or style_name.startswith("Heading"):
                # Flush previous section content before starting a new heading
                flush_current_section()
                current_section = text
            else:
                current_paragraphs.append(text)

        # 2. Extract tables where practical
        table_texts: List[str] = []
        for table_idx, table in enumerate(doc.tables):
            rows_data = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                if any(row_cells):
                    rows_data.append(" | ".join(row_cells))
            if rows_data:
                table_representation = f"\n[Table {table_idx + 1}]\n" + "\n".join(rows_data)
                table_texts.append(table_representation)

        if table_texts:
            current_paragraphs.extend(table_texts)

        # Flush final section
        flush_current_section()

        logger.info(f"Successfully extracted {len(sections)} sections from DOCX '{filename}'")
        return sections
